"""Экспорт заказов в форматы, которые принимает 1С (CSV с «;» и UTF-8 BOM, XLSX)."""

import csv
import re
import uuid
from datetime import datetime
from itertools import groupby
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from app.schemas import OrderRecommendation

COLUMNS = ["Поставщик", "Код 1С", "Артикул поставщика", "Наименование", "Ед.", "Количество", "Рекомендовано",
           "Статус", "Срочность", "Комментарий"]
_STATUS = {"pending": "На проверке", "approved": "Утверждено", "rejected": "Отклонено"}
_URGENCY = {"high": "Высокая", "medium": "Средняя", "low": "Низкая"}
FILENAME_RE = re.compile(r"^order-\d{8}-\d{6}-[0-9a-f]{6}\.(csv|xlsx|pdf)$")


def _qty(r: OrderRecommendation) -> float:
    """В заказ идёт утверждённое количество; для неутверждённых — рекомендованное."""
    return r.approved_qty if r.approved_qty is not None else r.recommended_qty


def _row(r: OrderRecommendation) -> list:
    return [r.supplier_name, r.sku_code, r.supplier_sku, r.name, r.unit, _qty(r), r.recommended_qty,
            _STATUS[r.status], _URGENCY[r.urgency], r.comment or ""]


def export(items: list[OrderRecommendation], fmt: str, export_dir: Path) -> str:
    export_dir.mkdir(parents=True, exist_ok=True)
    items = sorted((i for i in items if i.status != "rejected"), key=lambda r: (r.supplier_name, r.name))
    filename = f"order-{datetime.now():%Y%m%d-%H%M%S}-{uuid.uuid4().hex[:6]}.{fmt}"
    path = export_dir / filename

    if fmt == "csv":
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(COLUMNS)
            for r in items:
                w.writerow([f"{v:g}".replace(".", ",") if isinstance(v, float) else v for v in _row(r)])
    elif fmt == "xlsx":
        wb = Workbook()
        wb.remove(wb.active)
        for supplier, group in groupby(items, key=lambda r: r.supplier_name):
            ws = wb.create_sheet(title=supplier[:31])
            ws.append(COLUMNS)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in group:
                ws.append(_row(r))
            ws.column_dimensions["D"].width = 60
            for col in "BCE":
                ws.column_dimensions[col].width = 18
        if not wb.sheetnames:
            wb.create_sheet("Заказ").append(COLUMNS)
        wb.save(path)
    else:
        raise NotImplementedError("Экспорт в PDF пока не реализован — используйте csv или xlsx")
    return filename
