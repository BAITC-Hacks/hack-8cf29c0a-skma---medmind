"""Bounded CSV/XLSX parsing and atomic, repeatable imports into input tables."""

import csv
import io
import json
import zipfile
from datetime import date, datetime
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models import ImportIdentity
from app.services import input_data as data

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000


class ImportOptions(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mapping: dict[str, str] = Field(default_factory=dict, description="Заголовок файла → поле API")
    defaults: dict = Field(default_factory=dict)
    source: str = Field(default="1c", min_length=1, max_length=100, pattern=r"\S")
    sheet: str | None = None
    header_row: int = Field(default=1, ge=1, le=100)
    encoding: str = Field(default="utf-8-sig", pattern=r"^(utf-8-sig|utf-8|cp1251)$")
    delimiter: str = Field(default=";", pattern=r"^[;,\t]$")


def check_zip(blob: bytes, max_entries=200, max_size=100 * 1024 * 1024):
    try:
        archive = zipfile.ZipFile(io.BytesIO(blob))
        entries = archive.infolist()
        if len(entries) > max_entries or sum(i.file_size for i in entries) > max_size:
            archive.close()
            raise HTTPException(413, "Превышен лимит распакованного архива")
        if any(i.flag_bits & 1 for i in entries):
            archive.close()
            raise HTTPException(422, "Зашифрованные архивы не поддерживаются")
        return archive
    except zipfile.BadZipFile as exc:
        raise HTTPException(422, "Повреждённый ZIP/XLSX") from exc


def read_rows(blob: bytes, filename: str, options: ImportOptions):
    if len(blob) > MAX_BYTES:
        raise HTTPException(413, "Максимальный размер CSV/XLSX — 20 МБ")
    suffix = Path(filename).suffix.lower()
    workbook = None
    try:
        if suffix == ".csv":
            rows = csv.reader(io.StringIO(blob.decode(options.encoding)), delimiter=options.delimiter, strict=True)
        elif suffix == ".xlsx":
            with check_zip(blob):
                pass
            workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=False)
            sheet = workbook[options.sheet] if options.sheet else workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
        else:
            raise HTTPException(415, "Поддерживаются CSV и XLSX")
        for _ in range(options.header_row - 1):
            next(rows, None)
        header = next(rows, None)
        if not header:
            raise HTTPException(422, "Файл не содержит заголовков")
        headers = [str(v).strip() if v is not None else "" for v in header]
        # Excel may include empty formatted columns at the right edge.
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not h for h in headers) or len(set(headers)) != len(headers) or len(headers) > 100:
            raise HTTPException(422, "Заголовки должны быть непустыми и уникальными, максимум 100 колонок")
        if set(options.mapping) - set(headers):
            raise HTTPException(422, "Колонки из mapping отсутствуют в файле")
        result = []
        for row_number, row in enumerate(rows, options.header_row + 1):
            if row_number > options.header_row + MAX_ROWS:
                raise HTTPException(413, "Максимум 50 000 строк; разделите выгрузку на части")
            if not any(v is not None and str(v).strip() for v in row):
                continue
            if any(v is not None and str(v).strip() for v in row[len(headers):]):
                raise HTTPException(422, f"Строка {row_number}: больше значений, чем заголовков")
            result.append((row_number, dict(zip(headers, row, strict=False))))
        if not result:
            raise HTTPException(422, "Нет строк данных")
        return headers, result
    except HTTPException:
        raise
    except Exception as exc:
        # XML/ZIP/encoding engines use different exception classes for malformed uploads.
        raise HTTPException(422, "Не удалось прочитать файл: проверьте кодировку, лист и структуру") from exc
    finally:
        if workbook:
            workbook.close()


def convert(field: str, value):
    if value is None or value == "":
        return None
    if isinstance(value, str):
        value = value.strip()
        if value.startswith("="):
            raise ValueError(f"{field}: формулы не допускаются, выгрузите значения")
    if field in {"code", "sku_code", "id", "supplier_id", "category_id", "supplier_sku", "external_id", "document"}:
        if not isinstance(value, str):
            raise ValueError(f"{field}: идентификатор должен быть текстом (для сохранения ведущих нулей)")
    if field in {"qty", "unit_cost", "on_hand", "reserved", "min_order_qty", "order_multiple", "revenue",
                 "buffer_multiplier", "growth_coef", "seasonality_coef", "cover_months", "order_qty"}:
        if isinstance(value, str):
            value = value.replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if field in {"ts", "month", "as_of", "expected_date"}:
        if isinstance(value, datetime) and field != "ts":
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, str) and "." in value:
            for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y"):
                try:
                    parsed = datetime.strptime(value, fmt)
                    return parsed if field == "ts" else parsed.date()
                except ValueError:
                    continue
    if field == "locations" and isinstance(value, str):
        return json.loads(value)
    return value


def import_table(db: Session, name: str, blob: bytes, filename: str, options: ImportOptions, dry_run: bool):
    spec = data.resource(name)
    headers, rows = read_rows(blob, filename, options)
    mapping = options.mapping or {h: h for h in headers}
    allowed = set(spec.schema.model_fields) | ({"external_id"} if spec.generated else set())
    if set(mapping.values()) - allowed or set(options.defaults) - allowed:
        raise HTTPException(422, {"message": "Неизвестные поля", "allowed": sorted(allowed)})
    if len(set(mapping.values())) != len(mapping):
        raise HTTPException(422, "Несколько колонок сопоставлены одному полю")
    data.begin_write(db)
    counts = {"created": 0, "updated": 0, "unchanged": 0}
    errors, seen = [], set()
    for line, row in rows:
        try:
            with db.begin_nested():
                payload = dict(options.defaults)
                for header, field in mapping.items():
                    value = row.get(header)
                    if value is not None and str(value).strip():
                        payload[field] = convert(field, value)
                external_id = payload.pop("external_id", None)
                if spec.generated and (not isinstance(external_id, str) or not external_id.strip()
                                       or len(external_id) > 200):
                    raise ValueError("Для продаж и поставок обязателен текстовый external_id строки из 1С")
                values = data.validate(spec, payload)
                key = external_id if spec.generated else data.record_key(spec, values)
                if key in seen:
                    raise ValueError("Повтор ключа в одном файле")
                seen.add(key)
                identity = db.get(ImportIdentity, (name, options.source, external_id)) if spec.generated else None
                existing = data.find(db, spec, identity.record_key) if identity else (
                    None if spec.generated else data.find(db, spec, key))
                before = data.dump(existing) if existing else None
                record = data.save(db, name, values, existing, source=options.source)
                after = data.dump(record)
                counts["unchanged" if before == after else "updated" if before else "created"] += 1
                if spec.generated:
                    if identity is None:
                        identity = ImportIdentity(resource=name, source=options.source, external_id=external_id)
                    identity.record_key = data.record_key(spec, after)
                    db.add(identity)
        except (HTTPException, ValueError, IntegrityError) as exc:
            errors.append({"row": line, "detail": exc.detail if isinstance(exc, HTTPException) else
                           "Конфликт уникальности записи" if isinstance(exc, IntegrityError) else str(exc)})
            if len(errors) >= 100:
                break
    report = {"resource": name, "dry_run": dry_run, "applied": not dry_run and not errors,
              "total_rows": len(rows), **counts, "errors": errors, "errors_truncated": len(errors) >= 100}
    if dry_run or errors:
        db.rollback()
    else:
        db.commit()
    return report
